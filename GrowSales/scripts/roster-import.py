#!/usr/bin/env python3
# Parse the "Blitz Sign Up" workbook into the live weekly roster.
#
# The sheet is a grid: row 1 holds Sunday dates (one column per week), column A
# holds region anchors ("Houston (40 - 50 Reps)"), and each week-column stacks
# team blocks — a label or blank row separates blocks, the FIRST kept name in a
# block is the manager. Formatting carries meaning:
#   fill FF00FF (magenta) -> subcontractor       -> bucketed under "Subs"
#   fill FF0000 (red)     -> signed up, bailed   -> dropped entirely
# Everything else is region-band tinting and means nothing.
#
#   python3 scripts/roster-import.py ~/Downloads/"Blitz Sign Up-6.xlsx"            # preview current week
#   python3 scripts/roster-import.py file.xlsx --week 2026-07-26 --write           # push to the live roster
#
# --write upserts report_config.__roster_current__, which the dashboard AND the
# weekly-email API read on top of the deployed team-*.json (no deploy needed).

import sys, json, re, datetime, urllib.request, ssl
import openpyxl

SB = 'https://cmkfohwbejcztctrpdnt.supabase.co'
ANON = 'eyJhbGciOiJIUzI1NiIsInR5cCI6IkpXVCJ9.eyJpc3MiOiJzdXBhYmFzZSIsInJlZiI6ImNta2ZvaHdiZWpjenRjdHJwZG50Iiwicm9sZSI6ImFub24iLCJpYXQiOjE3ODIyNDYyODgsImV4cCI6MjA5NzgyMjI4OH0.waLa5E9a6SMvJ70keY9duAz4VxK39sDNNbmPSSbcebo'

SUB_FILL = 'FF00FF'
BAIL_FILL = 'FF0000'
STATUS = re.compile(r'^(summer|blitz|closed)\b', re.I)
# words that mark a cell as a location/note label, never a person
LABELISH = re.compile(r'\b(paying|rent|lyte|fiber|route|tbd|need|house|housing)\b', re.I)

# ---- name cleaning: mirrors scripts/_name-clean.mjs ----
SPELLING = {
  # NOTE: the deals feed and view picker spell him "Lagrange", so the roster
  # key must too (differs from _name-clean.mjs on purpose).
  'syndey turner': 'Sydney Turner', 'chris lagrange': 'Chris Lagrange',
  'christopher lagrange': 'Chris Lagrange', 'chris legrange': 'Chris Lagrange',
  'macee mofatt': 'Macee Moffat', 'calvin astil': 'Calvin Astill',
  'samuel hossfi': 'Samuel Hossfield', 'ethan hendrickson': 'Ethan Hendriksen',
  'gaberiel robbins': 'Gabriel Robbins', 'spencer schwieger': 'Spencer Schweiger',
  'ali ai ghazzawi': 'Ali Al Ghazzawi',
}
DOW = r'(?:sun|mon|tues?|wed|thur?s?|fri|sat)'
def clean_name(raw):
    s = str(raw or '')
    s = re.sub(r'\([^)]*\)', ' ', s)
    s = re.sub(r'\([^)]*$', ' ', s)
    s = re.sub(r'\$\S*', ' ', s)
    s = re.sub(rf'\b{DOW}\s*[-–]\s*{DOW}\b', ' ', s, flags=re.I)
    s = re.sub(rf'\b{DOW}\b', ' ', s, flags=re.I)
    s = re.sub(r'[0-9]+', ' ', s)
    s = re.sub(r'\b(LT|BNB|IN\s*MARKET|INMARKET|remotely|good|car)\b', ' ', s, flags=re.I)
    s = re.sub(r'[-–|].*$', ' ', s)
    s = re.sub(r'\s+', ' ', s).strip()
    s = re.sub(r'(?:\s+[A-Z])+$', '', s).strip()
    if not s: return ''
    s = SPELLING.get(s.lower(), s)
    return ' '.join(w[0].upper()+w[1:] if w and w == w.lower() else w for w in s.split(' '))

def fill_of(cell):
    f = cell.fill
    if f and f.fill_type == 'solid':
        rgb = f.start_color.rgb
        if isinstance(rgb, str): return rgb[-6:]
    return None

def parse_week(path, week_iso, tab=None):
    wb = openpyxl.load_workbook(path)
    target = datetime.date.fromisoformat(week_iso)
    ws = col = None
    for name in ([tab] if tab else wb.sheetnames):
        w = wb[name]
        for c in range(2, w.max_column + 1):
            v = w.cell(row=1, column=c).value
            if isinstance(v, datetime.datetime) and v.date() == target:
                ws, col = w, c; break
        if ws: break
    if not ws:
        sys.exit(f'No column headed {week_iso} found in any tab.')

    teams, subs, bailed, oddities = {}, [], [], []
    order, current_mgr, in_block = [], None, False
    for r in range(2, ws.max_row + 1):
        cell = ws.cell(row=r, column=col)
        raw = cell.value
        txt = str(raw).strip() if raw is not None else ''
        region = ws.cell(row=r, column=1).value            # region anchor resets blocks
        if (region and str(region).strip()) or not txt:
            in_block = False; continue
        if STATUS.match(txt) or LABELISH.search(txt):
            in_block = False; continue
        fl = fill_of(cell)
        name = clean_name(txt)
        if fl == BAIL_FILL:
            if name: bailed.append(name)
            continue                                        # a bailed row keeps the block open
        # Sheet convention: people are TitleCase. A fully-lowercase word in the
        # raw text ("Sherwood park") marks a place label, not a person.
        stripped = re.sub(r'\([^)]*\)?', ' ', txt)
        stripped = re.sub(rf'\b{DOW}\s*[-–]\s*{DOW}\b|\b{DOW}\b', ' ', stripped, flags=re.I)
        stripped = re.sub(r'\b(LT|BNB)\b|[A-Z]$', ' ', stripped)
        lowercase_word = re.search(r'\b[a-z]{3,}\b', stripped)
        if len(name.split(' ')) < 2 or '/' in txt or lowercase_word:
            in_block = False                                # market label (Baytown, Uvalde…)
            continue
        if fl == SUB_FILL:
            subs.append(name); continue                     # subs bucket, never a team lead
        if not in_block:
            current_mgr, in_block = name, True
            if name not in teams: teams[name] = []; order.append(name)
        elif name != current_mgr:
            teams[current_mgr].append(name)
    if subs:
        teams['Subs'] = sorted(set(subs))
    return ws.title, teams, order, sorted(set(bailed)), oddities

def main():
    args = sys.argv[1:]
    if not args: sys.exit(__doc__ or 'usage: roster-import.py <xlsx> [--week YYYY-MM-DD] [--tab NAME] [--write]')
    path = args[0]
    def opt(flag, default=None):
        return args[args.index(flag)+1] if flag in args else default
    today = datetime.date.today()
    sunday = today - datetime.timedelta(days=(today.weekday()+1) % 7)
    week = opt('--week', sunday.isoformat())
    tab, teams, order, bailed, _ = parse_week(path, week, opt('--tab'))

    total = sum(len(v) for v in teams.values()) + len([m for m in order])
    print(f'Week {week} (tab "{tab}"): {len(order)} teams, {total} people'
          f'{", "+str(len(teams.get("Subs",[])))+" subs" if teams.get("Subs") else ""}'
          f'{", "+str(len(bailed))+" bailed (dropped)" if bailed else ""}\n')
    for m in order:
        print(f'  {m} ({len(teams[m])}): {", ".join(teams[m]) or "—"}')
    if teams.get('Subs'): print(f'  Subs ({len(teams["Subs"])}): {", ".join(teams["Subs"])}')
    if bailed: print(f'\n  Bailed, dropped: {", ".join(bailed)}')

    if '--write' not in args:
        print('\nPreview only. Add --write to push to the live roster.'); return
    body = json.dumps({'manager': '__roster_current__',
                       'config': {'week': week, 'teams': teams,
                                  'saved_at': datetime.datetime.utcnow().isoformat()+'Z',
                                  'source': 'roster-import.py'}}).encode()
    req = urllib.request.Request(
        SB + '/rest/v1/report_config?on_conflict=manager', data=body, method='POST',
        headers={'apikey': ANON, 'Authorization': 'Bearer '+ANON,
                 'Content-Type': 'application/json',
                 'Prefer': 'resolution=merge-duplicates'})
    ctx = ssl.create_default_context(); ctx.check_hostname = False; ctx.verify_mode = ssl.CERT_NONE
    with urllib.request.urlopen(req, context=ctx) as resp:
        print(f'\nPushed to live roster ({resp.status}). Dashboard and email API pick it up on next load.')

if __name__ == '__main__':
    main()

#!/usr/bin/env python3
"""Lyte's serving areas for the Lyte map (More > ISP Maps).

    python3 scripts/lyte-areas.py "~/Downloads/GROW Headcount File - Living Doc.xlsx"
    python3 scripts/lyte-areas.py            # re-read the drops, keep the last turf sheet

Lyte runs everything by serving area (FDH): TXBYTWAC010.SA05, said out loud as
"Baytown 10.05". This rebuilds data/sectors/Lyte.csv (street + zip -> area) from
every file in Lyte/Raw on the drive, reading the area from, in order:
  the row's own "Serving Area Name" / "FDH" column,
  the sheet name ("Frost 1.4 Live", "Kingsville 1.2"),
  the file name, when it names exactly one area,
then fills in doors from SalesRabbit import batches whose file names one area
("SR LY-BAY (10.01) - Import.csv").

The headcount file's turf tracker (FDH, market, HHU, go live, status, last worked,
pass, bucket, last vendor) goes to data/sectors/Lyte-turf.json, keyed by the same
area name, and the map shows it on each area as Lyte's sheet has it.

Shorthand is only trusted when the FDH it spells is on Lyte's tracker, so a date
in a sheet name ("go live 9.14") can never invent an area.
Then: python3 scripts/sector-maps.py Lyte
"""
import collections, csv, datetime, glob, json, os, re, shutil, sys

import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(ROOT, 'scripts'))
import importlib.util
_spec = importlib.util.spec_from_file_location('si', os.path.join(ROOT, 'scripts', 'sector-import.py'))
si = importlib.util.module_from_spec(_spec); _spec.loader.exec_module(si)

SECTORS = os.path.join(ROOT, 'data', 'sectors')
OUT = os.path.join(SECTORS, 'Lyte.csv')
TURF = os.path.join(SECTORS, 'Lyte-turf.json')
RAW = si.raw_dir('Lyte')

FDH_RE = re.compile(r'^TX([A-Z]{4})AC(\d{3})\.SA(\d{2})([A-Z0-9]*)$')
# towns whose code no drop file has spelled out yet (Navarro county's tracker rows all say
# "Navarro County"); a code missing here and from the drops shows as the code itself
TOWN_FIX = {'KRNS': 'Kerens', 'POWL': 'Powell', 'MARN': 'Marion County', 'CRTH': 'Carthage',
            'CRCY': 'Crystal City'}
TURF_COLS = ['FDH', 'Market', 'City', 'SFU', 'Live_HHU', 'Forecasted_Go_Live', 'Actual_Go_Live', 'Primary_Status',
             'Last_Worked', 'Pass_Number', 'Turf Cycle', 'Days_Since_Worked', 'Bucket', 'Vendor_Last', 'Priority', 'Notes']


def cell(v):
    if v is None:
        return ''
    if isinstance(v, datetime.datetime):
        return v.date().isoformat()
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


# ---------------------------------------------------------------- turf tracker

def read_turf(path):
    """Mike's turf tracker tab: one row per FDH, Lyte's own read of every area."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = next((w for w in wb.worksheets if 'TURF TRACKER' in w.title.upper() and 'MANUAL' not in w.title.upper()), None)
    if ws is None:
        sys.exit(f'no turf tracker tab in {path}')
    rows = ws.iter_rows(values_only=True)
    hdr = [cell(x) for x in next(rows)]
    out = {}
    for r in rows:
        r = dict(zip(hdr, [cell(x) for x in r]))
        if FDH_RE.match(r.get('FDH', '')):
            out[r['FDH']] = {k: r.get(k, '') for k in TURF_COLS}
    # the tab is named for the day it was cut ("MIKES TURF TRACKER - 6.9")
    m = re.search(r'(\d{1,2})\.(\d{1,2})\s*$', ws.title)
    asof = ''
    if m:
        asof = f'{datetime.date.today().year}-{int(m.group(1)):02d}-{int(m.group(2)):02d}'
        if asof > datetime.date.today().isoformat():
            asof = f'{datetime.date.today().year - 1}' + asof[4:]
    print(f'turf tracker "{ws.title}": {len(out)} areas')
    return out, asof, ws.title


# ---------------------------------------------------------------- drop files

def sheets(path):
    """(sheet name, rows as lists) for every sheet of an xlsx, or the one csv."""
    if path.lower().endswith('.csv'):
        with open(path, newline='', encoding='utf-8-sig', errors='replace') as f:
            yield '', [[c.strip() for c in r] for r in csv.reader(f)]
        return
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    for ws in wb.worksheets:
        yield ws.title, [[cell(v) for v in r] for r in ws.iter_rows(values_only=True)]


def shorthand(text):
    """Every "10.05" / "1.2B" / "2.04A1" in a sheet or file name, dates and go-live notes removed."""
    t = re.sub(r'\d{1,2}\.\d{1,2}\.\d{2,4}', ' ', text)            # 9.17.2026
    t = re.sub(r'(?i)(go ?live|cip|live)\b.*$', ' ', t)               # "go live 9.14", "CIP Oct8"
    t = re.sub(r'\(\s*\d{1,2}[ _]\d{1,2}\s*\)|\b\d{1,2}_\d{1,2}\b', ' ', t)   # (5_27), 7_8
    return [(int(h), int(s), (x or '').upper()) for h, s, x in
            re.findall(r'(?<![\d.])(\d{1,2})\.(\d{1,2})\s?([A-Za-z]\d?)?(?![\d.])', t)]


def spell(code, h, s, x):
    return f'TX{code}AC{h:03d}.SA{s:02d}{x}'


def resolve(code, sh, known):
    """An FDH from shorthand, only if Lyte's tracker has it (else the bare-letter form)."""
    h, s, x = sh
    for cand in (spell(code, h, s, x), spell(code, h, s, x[:1]), spell(code, h, s, '')):
        if cand in known:
            return cand
    return None


def find(hdr, *pats):
    for p in pats:
        for i, h in enumerate(hdr):
            if re.search(p, h, re.I):
                return i
    return None


def read_drop(path, known, code_of_city, town_hint):
    """Rows (street_key, zip, fdh, city, state, how) out of one Lyte drop file."""
    base = os.path.basename(path)
    file_sh = shorthand(re.sub(r'\.(xlsx|csv)$', '', base))
    out, notes = [], []
    for title, rows in sheets(path):
        rows = [r for r in rows if any(r)]
        if not rows:
            continue
        hdr = [h.strip() for h in rows[0]]
        if hdr and re.match(r'^\d+$', hdr[0]):          # Kingsville files: no header row
            hdr, body = ['Number', 'Street', 'City', 'State', 'Zip'], rows
        else:
            body = rows[1:]
        c_fdh = find(hdr, r'^serving ?area', r'^fdh$')
        c_addr = find(hdr, r'^address$')
        c_num = find(hdr, r'^street ?number$', r'^number$')
        c_st = find(hdr, r'^street$')
        c_zip = find(hdr, r'^zip', r'postal')
        c_city = find(hdr, r'^city$')
        c_state = find(hdr, r'^state')
        c_status = find(hdr, r'^status$')
        if c_addr is None and c_st is None:
            notes.append(f'{title or "csv"}: no address column'); continue
        sheet_sh = shorthand(title)
        n = 0
        for r in body:
            g = lambda i: (r[i] if i is not None and i < len(r) else '').strip()
            st = g(c_addr) if c_addr is not None else f'{g(c_num)} {g(c_st)}'.strip()
            if not st or 'UNKNOWN ADDRESS' in st.upper():
                continue
            city = g(c_city).title()
            fdh = g(c_fdh).upper() if c_fdh is not None else ''
            how = 'column'
            if not FDH_RE.match(fdh):
                fdh, code = '', code_of_city.get(city.upper()) or town_hint
                for src, sh in (('sheet', sheet_sh), ('file', file_sh)):
                    if code and len(sh) == 1:
                        fdh = resolve(code, sh[0], known) or ''
                        if fdh:
                            how = src; break
            if not fdh:
                continue
            out.append((si.street_key(st), re.sub(r'\D', '', g(c_zip))[:5], fdh, city, g(c_state) or 'TX', how))
            n += 1
        notes.append(f'{title or "csv"}: {n}')
    return out, notes


# ---------------------------------------------------------------- SalesRabbit batches

def batch_rows(known, code_of_city):
    """Doors from SalesRabbit import batches whose file name names exactly one area."""
    reg = json.load(open(os.path.join(ROOT, 'data', 'isp-imports.json')))
    one = {}
    for imp, v in reg.items():
        if v.get('isp') != 'Lyte' or not v.get('file'):
            continue
        sh = set(shorthand(re.sub(r'\.csv$', '', v['file'])))
        if len(sh) == 1:
            one[imp] = sh.pop()
    doors = {d['id']: d for d in json.load(open(os.path.join(ROOT, 'data', 'osm-cache', 'Lyte-doors.json')))}
    imp_of = {}
    for line in open(os.path.join(ROOT, 'data', 'sr-leads.jsonl')):
        if '"import_id"' not in line:
            continue
        try:
            o = json.loads(line)
        except ValueError:
            continue
        if o['id'] in doors and o.get('import_id') in one:
            imp_of[o['id']] = o['import_id']
    out, missed = [], collections.Counter()
    for i, imp in imp_of.items():
        d = doors[i]
        code = code_of_city.get((d['city'] or '').upper())
        fdh = resolve(code, one[imp], known) if code else None
        if fdh:
            out.append((si.street_key(d['street']), d['zip'][:5], fdh, d['city'], d['state'] or 'TX', 'batch'))
        else:
            missed[reg[imp]['file']] += 1
    return out, missed


# ---------------------------------------------------------------- names

def area_names(fdhs, town_of):
    """TXBYTWAC010.SA05 -> "Baytown 10.05", the way Lyte's sheets say it."""
    out = {}
    for f in fdhs:
        code, h, s, x = FDH_RE.match(f).groups()
        out[f] = f'{town_of.get(code, code.title())} {int(h)}.{s}{x}'
    return out


def main():
    args = [os.path.expanduser(a) for a in sys.argv[1:] if not a.startswith('--')]
    if args:
        turf, asof, tab = read_turf(args[0])
        dest = os.path.join(RAW, datetime.date.today().isoformat() + ' ' + os.path.basename(args[0]))
        if os.path.isdir(RAW) and not os.path.exists(dest):
            shutil.copy2(args[0], dest); print(f'  archived -> {dest}')
        src = os.path.basename(args[0])
    else:
        prev = json.load(open(TURF))
        turf = {v['FDH']: {k: v.get(k, '') for k in TURF_COLS} for v in prev['areas'].values()}
        asof, tab, src = prev['asof'], prev['tab'], prev['source']
    known = set(turf)

    # which town each 4-letter code is, and which code a city's doors use: from every
    # drop row that carries its FDH, then the tracker's own City column
    files = sorted(p for p in glob.glob(os.path.join(RAW, '*')) if p.lower().endswith(('.xlsx', '.csv'))
                   and not re.search(r'dupes|SalesRabbit status|Headcount', os.path.basename(p), re.I))
    city_code = collections.defaultdict(collections.Counter)
    for p in files:
        rows, _ = read_drop(p, set(), {}, None)      # column rows only
        for _, _, f, city, _, how in rows:
            if how == 'column' and city:
                city_code[city.upper()][FDH_RE.match(f).group(1)] += 1
    for f, t in turf.items():
        c = (t['City'] or '').upper()
        if c and 'COUNTY' not in c:
            city_code[c][FDH_RE.match(f).group(1)] += 1
    code_of_city = {c: n.most_common(1)[0][0] for c, n in city_code.items()}
    # a code is its town's letters (RCLD = Richland, STMN = Streetman): a Navarro county
    # file lists Corsicana on Richland's doors too, so the plain majority would name
    # two codes the same town
    def spells(code, city):
        it = iter(re.sub(r'[^A-Z]', '', city))
        return city[:1] == code[:1] and all(ch in it for ch in code)
    town_of = {}
    for c, n in city_code.items():
        for code, k in n.items():
            fit = spells(code, c)
            if (fit, k) > town_of.get(code, ('', False, 0))[1:]:
                town_of[code] = (c.title(), fit, k)
    town_of = {k: v[0] for k, v in town_of.items() if v[1]}
    # codes no drop has spelled out yet: Lyte's tracker only says "Navarro County"
    for code, town in TOWN_FIX.items():
        town_of.setdefault(code, town)

    rows = {}
    for p in files:
        hint = None
        m = re.search(r'(Kingsville|Marshall|Baytown|Uvalde|Beeville|Atlanta)', os.path.basename(p), re.I)
        if m:
            hint = code_of_city.get(m.group(1).upper())
        got, notes = read_drop(p, known, code_of_city, hint)
        for r in got:                       # newest file last, so it wins
            rows[(r[0], r[1])] = r
        print(f'  {os.path.basename(p)[:60]:<60} {sum(1 for _ in got):>5}  ' + ', '.join(notes)[:90])
    filled, missed = batch_rows(known, code_of_city)
    added = 0
    for r in filled:
        if (r[0], r[1]) not in rows:
            rows[(r[0], r[1])] = r; added += 1
    print(f'  SalesRabbit batches named for one area: {added:,} more doors')
    for f, n in missed.most_common(8):
        print(f'    no area on the tracker for {n} doors of {f}')

    names = area_names({r[2] for r in rows.values()} | known, town_of)
    for p in (OUT, TURF):
        if os.path.exists(p) and not os.path.exists(p + '.bak-pre-lyte-areas'):
            shutil.copy2(p, p + '.bak-pre-lyte-areas')
    with open(OUT, 'w', newline='') as f:
        w = csv.writer(f)
        w.writerow(['street_key', 'zip', 'sector', 'city', 'state', 'fdh'])
        for k, z, fdh, city, state, _ in sorted(rows.values(), key=lambda r: (names[r[2]], r[0])):
            w.writerow([k, z, names[fdh], city, state, fdh])
    json.dump({'asof': asof, 'tab': tab, 'source': src,
               'areas': {names[f]: t for f, t in turf.items()}}, open(TURF, 'w'), indent=0)
    how = collections.Counter(r[5] for r in rows.values())
    print(f'\n{OUT}: {len(rows):,} addresses in {len({r[2] for r in rows.values()})} areas ({dict(how)})')
    print(f'{TURF}: {len(turf)} areas from "{tab}"')


if __name__ == '__main__':
    main()

#!/usr/bin/env python3
"""File side of the Lead desk: read an ISP workbook, or rewrite it in the one
shape lead-import.py always understands. No guessing happens here; the column
mapping comes from workers/grow/leaddesk.mjs (a saved format or a Haiku call).

    leaddesk.py inspect   <file>                      -> JSON: sheets, first rows, row counts
    leaddesk.py normalize <file> <mapping.json> <out> -> writes <out> (xlsx), JSON counts per sheet

mapping.json: {"sheets": {"<sheet name>": {"header_row": 0, "cols": {"street": 3, "zip": 7, ...}}}}
Column values are 0-based indexes. A sheet left out of the mapping is skipped.
"""
import csv, json, re, sys

csv.field_size_limit(10 ** 7)
OUT_HDR = ['Address', 'Street Number', 'Unit', 'City', 'State', 'Zip', 'Latitude', 'Longitude',
           'Sector', 'Status', 'Customer']
FIELDS = ['street', 'streetnum', 'unit', 'city', 'state', 'zip', 'lat', 'lng', 'sector', 'status', 'customer']
# "123 Main St, Glens Falls, NY 12801" in one cell
FULL = re.compile(r'^\s*(.+?),\s*([^,]+?),\s*([A-Za-z]{2})\.?,?\s+(\d{5})(?:-\d{4})?\s*(?:,.*)?$')


def cell(v):
    if v is None:
        return ''
    return re.sub(r'^(-?\d+)\.0$', r'\1', str(v).strip())


def sheets_of(path):
    """[(sheet name, rows)] with blank rows dropped. A csv is one sheet."""
    if path.lower().endswith(('.xlsx', '.xlsm')):
        import openpyxl
        wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
        out = []
        for ws in wb.worksheets:
            if ws.sheet_state != 'visible':
                continue
            rows = [[cell(v) for v in r] for r in ws.iter_rows(values_only=True)]
            out.append((ws.title, [r for r in rows if any(r)]))
        return out
    with open(path, newline='', encoding='utf-8-sig') as f:
        rows = [[c.strip() for c in r] for r in csv.reader(f)]
    return [('csv', [r for r in rows if any(r)])]


def inspect(path):
    out = []
    for name, rows in sheets_of(path):
        width = max((len(r) for r in rows[:50]), default=0)
        head = [(r + [''] * width)[:width] for r in rows[:12]]
        out.append({'name': name, 'rows': len(rows), 'width': width, 'head': head})
    print(json.dumps({'sheets': out}))


def pick(row, idx):
    return row[idx].strip() if isinstance(idx, int) and 0 <= idx < len(row) else ''


def normalize(path, mapping_path, out_path):
    import openpyxl
    mapping = json.load(open(mapping_path))['sheets']
    wb = openpyxl.Workbook(); ws = wb.active
    ws.append(OUT_HDR)
    report = []
    for name, rows in sheets_of(path):
        m = mapping.get(name)
        if not m:
            report.append({'name': name, 'skipped': True, 'rows': len(rows)})
            continue
        cols, body = m['cols'], rows[int(m.get('header_row', 0)) + 1:]
        n = split = 0
        for r in body:
            v = {f: pick(r, cols.get(f)) for f in FIELDS}
            if 'full' in cols and not (v['city'] and v['zip']):
                g = FULL.match(pick(r, cols['full']))
                if g:
                    split += 1
                    v['street'] = v['street'] or g.group(1)
                    v['city'] = v['city'] or g.group(2)
                    v['state'] = v['state'] or g.group(3).upper()
                    v['zip'] = v['zip'] or g.group(4)
                elif not v['street']:
                    v['street'] = pick(r, cols['full'])
            if not any(v[f] for f in ('street', 'streetnum', 'city', 'zip')):
                continue
            ws.append([v[f] for f in FIELDS])
            n += 1
        report.append({'name': name, 'rows': n, 'split_full_address': split})
    wb.save(out_path)
    print(json.dumps({'out': out_path, 'sheets': report,
                      'total': sum(s['rows'] for s in report if not s.get('skipped'))}))


if __name__ == '__main__':
    if len(sys.argv) == 3 and sys.argv[1] == 'inspect':
        inspect(sys.argv[2])
    elif len(sys.argv) == 5 and sys.argv[1] == 'normalize':
        normalize(sys.argv[2], sys.argv[3], sys.argv[4])
    else:
        sys.exit(__doc__)

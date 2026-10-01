#!/usr/bin/env python3
"""Render each slide of a .pptx to PNG using the Grow template geometry.

No LibreOffice or PowerPoint on this machine, so this reads the deck back with
python-pptx and draws it with PIL at 100 dpi. Arial Black stands in for Archivo
Black (similar weight and width) and Arial Bold for Garet, with a 12% width
penalty on body text because Garet runs wider than Arial. Anything that would
spill its box is outlined in red so overflow is visible at a glance.
"""
import sys, os, math
from pptx import Presentation
from pptx.util import Emu
from PIL import Image, ImageDraw, ImageFont

DPI = 100
W, H = int(13.333 * DPI), int(7.5 * DPI)
BG = (6, 9, 15)
FG = (255, 255, 255)
RED = (255, 80, 80)
TITLE_TTF = "/System/Library/Fonts/Supplemental/Arial Black.ttf"
BODY_TTF = "/System/Library/Fonts/Supplemental/Arial Bold.ttf"
BODY_PENALTY = 1.12

def px(inches): return int(inches * DPI)
def pt(points): return int(points * DPI / 72)

def wrap(draw, text, font, max_w, extra=0, penalty=1.0):
    words, lines, cur = text.split(), [], ""
    for w in words:
        t = (cur + " " + w).strip()
        width = (draw.textlength(t, font=font) + extra * len(t)) * penalty
        if width <= max_w or not cur:
            cur = t
        else:
            lines.append(cur); cur = w
    if cur: lines.append(cur)
    return lines

def render(pptx_path, out_dir):
    prs = Presentation(pptx_path)
    os.makedirs(out_dir, exist_ok=True)
    logo = None
    for i, slide in enumerate(prs.slides, 1):
        im = Image.new("RGB", (W, H), BG)
        d = ImageDraw.Draw(im)
        overflow = False
        for shp in slide.shapes:
            if shp.shape_type == 13 and logo is None:  # picture
                from io import BytesIO
                logo = Image.open(BytesIO(shp.image.blob)).convert("RGBA")
            if not shp.has_text_frame: continue
            x, y = px(Emu(shp.left).inches), px(Emu(shp.top).inches)
            w, h = px(Emu(shp.width).inches), px(Emu(shp.height).inches)
            tf = shp.text_frame
            paras = tf.paragraphs
            size = None
            for p in paras:
                for r in p.runs:
                    if r.font.size: size = r.font.size.pt; break
                if size: break
            size = size or 26
            is_title = size >= 44
            if is_title:
                font = ImageFont.truetype(TITLE_TTF, pt(size))
                text = "".join(r.text for p in paras for r in p.runs)
                spacing = pt(8) if size == 44 else pt(1)
                lines = wrap(d, text, font, w, extra=spacing)
                lh = int(pt(size) * 1.2)
                total = lh * len(lines)
                if size == 96:  # centered cover
                    ty = y + (h - total) // 2
                    for ln in lines:
                        tw = d.textlength(ln, font=font) + spacing * len(ln)
                        tx = x + (w - tw) // 2
                        cx = tx
                        for ch in ln:
                            d.text((cx, ty), ch, font=font, fill=FG)
                            cx += d.textlength(ch, font=font) + spacing
                        ty += lh
                else:
                    ty = y
                    for ln in lines:
                        cx = x
                        for ch in ln:
                            d.text((cx, ty), ch, font=font, fill=FG)
                            cx += d.textlength(ch, font=font) + spacing
                        ty += lh
                if total > h: overflow = True
            else:
                font = ImageFont.truetype(BODY_TTF, pt(size))
                lh = int(pt(size) * 1.2)
                gap = pt(10)
                ty = y + pt(4)
                for p in paras:
                    text = "".join(r.text for r in p.runs)
                    if not text: continue
                    lvl = p.level or 0
                    indent = px(0.45) + lvl * px(0.5)
                    bullet = "●" if lvl == 0 else "○"
                    d.text((x + lvl * px(0.5) + px(0.05), ty), bullet, font=font, fill=FG)
                    lines = wrap(d, text, font, w - indent, penalty=BODY_PENALTY)
                    for ln in lines:
                        d.text((x + indent, ty), ln, font=font, fill=FG)
                        ty += lh
                    ty += gap
                if ty - y > h: overflow = True
            if overflow:
                d.rectangle([x, y, x + w, y + h], outline=RED, width=4)
        if logo is not None:
            lw, lh_ = px(1.13), px(0.19)
            box = logo.copy(); box.thumbnail((lw, lh_))
            im.paste(box, (px(0.81), px(6.57)), box)
        name = os.path.join(out_dir, f"slide-{i:02d}.png")
        im.save(name)
        if overflow: print(f"  OVERFLOW slide {i}")
    return len(prs.slides)

if __name__ == "__main__":
    for p in sys.argv[1:]:
        base = os.path.splitext(os.path.basename(p))[0]
        out = os.path.join(os.path.dirname(p), "preview", base)
        n = render(p, out)
        print(f"{base}: {n} slides -> {out}")
